"""올린 파일에서 글을 뽑는다. 순수 함수: DB · 네트워크 없음.

표(엑셀 · CSV · 워드 표)는 행마다 `열이름: 값 · 열이름: 값` 한 줄로 만든다. 그래야 한 행이 한 덩어리 안에서
열 이름과 함께 읽힌다. 수식은 저장된 계산값만 읽고, 값이 없으면 빈칸으로 두고 개수를 알린다.
"""
import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from typing import Literal

SUPPORTED = {"xlsx", "csv", "txt", "md", "pdf", "docx"}
NO_TEXT_MESSAGE = "글자를 읽을 수 없는 파일입니다. 텍스트가 들어 있는 PDF나 문서로 올려 주세요."
MAX_CHARS = 1500
MAX_HEADING = MAX_CHARS // 2
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.?!。])\s+")


class ExtractError(Exception):
    """사용자에게 그대로 보여 줄 수 있는 한국어 이유를 담는다."""


@dataclass
class Block:
    kind: Literal["text", "table"]
    heading: str
    lines: list[str]


@dataclass
class Extracted:
    blocks: list[Block]
    notes: dict = field(default_factory=lambda: {"formula_without_value": 0, "truncated_rows": 0})


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractError("글자 인코딩을 알 수 없는 파일입니다. UTF-8로 저장해서 다시 올려 주세요.")


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "예" if v else "아니오"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def _split_sentences(text: str) -> list[str]:
    out: list[str] = []
    for raw in text.splitlines():
        for part in _SENTENCE_SPLIT_RE.split(raw):
            part = part.strip()
            if part:
                out.append(part)
    return out


def _table_block(title: str, rows: list[list[str]], notes: dict | None = None) -> Block | None:
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return None
    header, body = rows[0], rows[1:]
    names = [h or f"열{i + 1}" for i, h in enumerate(header)]
    lines = []
    for r in body:
        cells = [f"{names[i] if i < len(names) else f'열{i + 1}'}: {c}" for i, c in enumerate(r) if c]
        if cells:
            lines.append(" · ".join(cells))
    heading = " · ".join(p for p in (title, " | ".join(h for h in header if h)) if p)[:MAX_HEADING]
    budget = MAX_CHARS - len(heading) - 1  # 덩어리 한 개(머리글 + 행 1개)가 1,500자 안에 들어가야 한다
    cut = [ln[: budget - 1] if len(ln) + 1 > budget else ln for ln in lines]
    if notes is not None:
        notes["truncated_rows"] = notes.get("truncated_rows", 0) + sum(a != b for a, b in zip(lines, cut))
    return Block("table", heading, cut) if cut else None


def _text_blocks(text: str) -> list[Block]:
    """마크다운 `#` 제목이 나오면 새 블록으로 나눈다."""
    blocks: list[Block] = []
    heading, buf = "", []

    def flush() -> None:
        lines = _split_sentences("\n".join(buf))
        if lines:
            blocks.append(Block("text", heading, lines))

    for raw in text.splitlines():
        m = re.match(r"^#{1,6}\s+(.*\S)\s*$", raw)
        if m:
            flush()
            heading, buf = m.group(1), []
        else:
            buf.append(raw)
    flush()
    return blocks


def _xlsx(data: bytes) -> Extracted:
    from openpyxl import load_workbook

    try:
        wb_f = load_workbook(io.BytesIO(data), data_only=False)
        wb_v = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise ExtractError("엑셀 파일을 열 수 없습니다. 올바른 .xlsx 파일인지 확인해 주세요.") from e
    out = Extracted([])
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        rows: list[list[str]] = []
        for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            vals = []
            for cf, cv in zip(row_f, row_v):
                if cv.value is None and isinstance(cf.value, str) and cf.value.startswith("="):
                    out.notes["formula_without_value"] += 1
                vals.append(_cell_text(cv.value))
            rows.append(vals)
        block = _table_block(ws_f.title, rows, out.notes)
        if block:
            out.blocks.append(block)
    return out


def _csv(data: bytes) -> Extracted:
    try:
        rows = [[_cell_text(c) for c in r] for r in csv.reader(io.StringIO(_decode(data)))]
    except csv.Error as e:
        raise ExtractError("CSV 파일의 형식이 올바르지 않습니다. 쉼표로 구분된 표인지 확인해 주세요.") from e
    out = Extracted([])
    block = _table_block("", rows, out.notes)
    if block:
        out.blocks.append(block)
    return out


def _docx(data: bytes) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError("워드 파일을 열 수 없습니다. 올바른 .docx 파일인지 확인해 주세요.") from e
    out = Extracted([])
    blocks = out.blocks
    heading, buf = "", []

    def flush() -> None:
        lines = _split_sentences("\n".join(buf))
        if lines:
            blocks.append(Block("text", heading, lines))

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith(("Heading", "Title")) and p.text.strip():
                flush()
                heading, buf = p.text.strip(), []
            else:
                buf.append(p.text)
        elif tag == "tbl":
            flush()
            buf = []
            t = Table(child, doc)
            block = _table_block(heading, [[_cell_text(c.text) for c in r.cells] for r in t.rows], out.notes)
            if block:
                blocks.append(block)
    flush()
    return out


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractError("암호가 걸린 PDF입니다. 암호를 풀어서 다시 올려 주세요.")
        pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    except ExtractError:
        raise
    except Exception as e:
        raise ExtractError("PDF 파일을 열 수 없습니다. 올바른 PDF인지 확인해 주세요.") from e
    blocks = []
    for n, text in pages:
        lines = _split_sentences(text)
        if lines:
            blocks.append(Block("text", f"{n}쪽", lines))
    return Extracted(blocks)


def extract_text(file_type: str, data: bytes) -> Extracted:
    ft = (file_type or "").lower().lstrip(".")
    if ft not in SUPPORTED:
        raise ExtractError("지원하지 않는 파일 형식입니다. xlsx · csv · txt · md · pdf · docx 파일을 올려 주세요.")
    if ft == "xlsx":
        result = _xlsx(data)
    elif ft == "csv":
        result = _csv(data)
    elif ft == "docx":
        result = _docx(data)
    elif ft == "pdf":
        result = _pdf(data)
    else:
        result = Extracted(_text_blocks(_decode(data)))
    if not result.blocks:
        raise ExtractError(NO_TEXT_MESSAGE)
    return result
