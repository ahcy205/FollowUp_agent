import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.knowledge.chunking import chunk_blocks
from app.knowledge.extract import Block, ExtractError, extract_text
from app.mail.draft_checks import tokens


def _xlsx(rows, title="단가표"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(text: str | None) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n" % (len(objs) + 1, xref) + b"%%EOF"
    return out


def _num_keys(text):
    return [k for kind, k, _ in tokens(text) if kind == "n"]


def test_xlsx_rows_become_labeled_lines():
    ex = extract_text("xlsx", _xlsx([["품번", "단가", "최소수량"], ["TS-200", 12000, 100]]))
    b = ex.blocks[0]
    assert b.kind == "table"
    assert "단가표" in b.heading
    assert b.lines == ["품번: TS-200 · 단가: 12000 · 최소수량: 100"]


def test_xlsx_number_cells_render_without_float_suffix():
    ex = extract_text("xlsx", _xlsx([["a", "b"], [12000.0, 12.5]]))
    assert ex.blocks[0].lines == ["a: 12000 · b: 12.5"]
    assert _num_keys("12,000원") and _num_keys("12,000원") == _num_keys("a: 12000")


def test_xlsx_formula_without_cached_value_counted():
    ex = extract_text("xlsx", _xlsx([["a", "b"], [1, "=A2*2"]]))
    assert ex.notes["formula_without_value"] == 1
    assert ex.blocks[0].lines == ["a: 1"]


def test_csv_utf8_bom_and_cp949():
    text = "품번,단가\nTS-1,500\n"
    for data in (text.encode("utf-8-sig"), text.encode("cp949")):
        assert extract_text("csv", data).blocks[0].lines == ["품번: TS-1 · 단가: 500"]


def test_docx_paragraphs_and_table():
    d = Document()
    d.add_heading("납품 조건", level=1)
    d.add_paragraph("결제는 월말 마감 후 30일 이내입니다.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "품목", "단가"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "A", "100"
    d.add_heading("배송", level=1)
    d.add_paragraph("택배로 보냅니다.")
    buf = io.BytesIO()
    d.save(buf)
    blocks = extract_text("docx", buf.getvalue()).blocks
    assert [b.kind for b in blocks] == ["text", "table", "text"]
    assert blocks[0].heading == "납품 조건"
    assert blocks[1].lines == ["품목: A · 단가: 100"]
    assert blocks[2].heading == "배송"


def test_pdf_text_layer():
    assert any("Hello" in ln for b in extract_text("pdf", _pdf("Hello world")).blocks for ln in b.lines)
    with pytest.raises(ExtractError) as e:
        extract_text("pdf", _pdf(None))
    assert str(e.value) == "글자를 읽을 수 없는 파일입니다. 텍스트가 들어 있는 PDF나 문서로 올려 주세요."


def test_bad_files_raise_extract_error():
    for ft in ("pdf", "docx", "xlsx"):
        with pytest.raises(ExtractError):
            extract_text(ft, b"not a real file")


def test_table_chunks_keep_rows_whole_with_header():
    rows = [f"품번: P{i:03d} · 단가: {i * 100}" for i in range(45)]
    chunks = chunk_blocks([Block("table", "단가표 · 품번 | 단가", rows)])
    assert [len(c.lines) - 1 for c in chunks] == [20, 20, 5]
    assert all(c.lines[0] == "단가표 · 품번 | 단가" and len(c.text) <= 1500 for c in chunks)
    assert [ln for c in chunks for ln in c.lines[1:]] == rows
    long_rows = [f"메모: {'가' * 300} {i}" for i in range(10)]
    for c in chunk_blocks([Block("table", "표", long_rows)]):
        assert len(c.text) <= 1500 and c.lines[0] == "표"


def test_text_chunks_split_at_paragraph_with_overlap():
    lines = [f"{i:02d}번째 문단입니다. " + "가" * 180 for i in range(16)]  # 약 3,000자
    chunks = chunk_blocks([Block("text", "", lines)])
    assert len(chunks) >= 2
    assert all(len(c.text) <= 1500 for c in chunks)
    assert len(chunks[0].text) >= 1200
    assert chunks[0].lines[-1] in lines  # 줄 경계에서 닫힘
    for a, b in zip(chunks, chunks[1:]):
        assert b.text.startswith(a.text[-100:])
    assert chunks[-1].text.endswith(lines[-1])
